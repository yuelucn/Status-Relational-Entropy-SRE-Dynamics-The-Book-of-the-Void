# -*- coding: utf-8 -*-
"""验证标题中的公式是否真实存在（检查XML）"""
from docx import Document

p = r'C:\mywork\SRE-Dynamics\状态-关系熵（SRE）动力学_整理版.docx'
doc = Document(p)

import re
# 找"算子-1"标题段落
for pa in doc.paragraphs:
    if pa.text.startswith('算子-1：局部图扩张算子'):
        xml = pa._p.xml
        has_math = '<m:oMath' in xml
        # 提取公式文本
        math_texts = re.findall(r'<m:t[^>]*>([^<]*)</m:t>', xml)
        print('标题文本:', pa.text)
        print('包含OMML公式:', has_math)
        print('公式片段:', ''.join(math_texts)[:100])
        break

# 检查一个普通段落的公式
for pa in doc.paragraphs:
    if '普朗克长度' in pa.text or '涌现本体紫外边界' in pa.text:
        xml = pa._p.xml
        has_math = '<m:oMath' in xml
        print()
        print('正文段包含OMML公式:', has_math, '| 段落:', pa.text[:40])
        break

# 统计含公式的标题数
math_heads = 0
for pa in doc.paragraphs:
    if pa.style.name.startswith('Heading') and '<m:oMath' in pa._p.xml:
        math_heads += 1
print('含公式的标题数:', math_heads)
